"""Render the LLM-authored Marp deck to a native .pptx with python-pptx.

Why this exists
---------------
``slides.pptx`` was produced by shelling out to ``pandoc slides.md -o
slides.pptx``. That has two problems on the machines that need slides most:

* **It needs pandoc.** On a locked-down host without it, there were no slides
  at all -- not even the .pptx, which is the one slide format that needs no
  renderer to *view*.
* **It is unthemed.** Pandoc's built-in pptx template is plain white Calibri.
  The quest's ``slides.html`` / ``slides.pdf`` carry FI's identity (warm paper,
  deep teal, serif headings) via ``templates/slides/fi.css``; the .pptx did
  not, so the one artifact a colleague is most likely to open looked nothing
  like the others.

This module renders the same ``slides.md`` directly, applying the fi.css
design language as native PowerPoint shapes. Pure ``python-pptx`` -- no
external binary -- so it works wherever FI itself runs, and the output stays
fully editable, which a rendered PDF is not.

Design is lifted from ``templates/slides/fi.css`` so the deck matches its HTML
sibling: warm paper ground, a short teal kicker rule above each serif slide
title, uppercase sans eyebrows, teal list markers, and the wordmark footer.

Supported Marp conventions (the subset ``agents/slides.md`` instructs the
model to emit):

* ``<!-- _class: lead -->``          -- title / closing slide, inverted
* ``# H1`` / ``## H2`` / ``### H3``  -- thesis / slide title / eyebrow
* ``- `` bullets (one nesting level), ``**bold**``, ``*italic*``, ``_italic_``
* ``![bg right:40% fit](figures/x)`` -- figure right, text left
* ``![w:900](...)`` / ``![h:420](...)`` / bare ``![](...)`` -- figure slide
* ``> blockquote``                   -- pull quote
* fenced code blocks
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from generation._pptx_math import math_xml, plain_text, split_math

_log = logging.getLogger("frontier_insight.slides")

# ---- fi.css design tokens (keep in sync with templates/slides/fi.css) ----
INK = (0x16, 0x22, 0x2B)      # --ink      near-black slate
PAPER = (0xF8, 0xF7, 0xF3)    # --paper    warm off-white
MUTED = (0x65, 0x72, 0x7C)    # --muted
FAINT = (0x8A, 0x94, 0x9B)    # --faint
HAIR = (0xE4, 0xE2, 0xDA)     # --hair
ACCENT = (0x0E, 0x6E, 0x6B)   # --accent   deep teal, the single accent
ACCENT_2 = (0x0A, 0x4F, 0x4D)  # --accent-2 darker teal

SERIF = "Palatino Linotype"   # matches fi.css --serif first choice
SANS = "Segoe UI"
MONO = "Consolas"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.9


def _rgb(t: tuple[int, int, int]):
    from pptx.dml.color import RGBColor
    return RGBColor(*t)


# ---------------------------------------------------------------- parsing

_FRONTMATTER = re.compile(r"\A---\s*\n.*?\n---\s*\n", re.DOTALL)
_IMG = re.compile(r"!\[(?P<attrs>[^\]]*)\]\((?P<src>[^)]+)\)")
_BG_RIGHT = re.compile(r"bg\s+right(?::(?P<pct>\d+)%)?")
_W_ATTR = re.compile(r"\bw:(\d+)")
_H_ATTR = re.compile(r"\bh:(\d+)")
_LEAD = re.compile(r"<!--\s*_class:\s*lead\s*-->")


class Slide:
    """One parsed slide: headings, body blocks, and an optional figure."""

    def __init__(self) -> None:
        self.lead = False
        self.h1 = ""
        self.h2 = ""
        self.h3 = ""
        # Paragraphs and bullets in the slide's order: ("para", 0, text) or
        # ("bullet", indent_level, text). A line under a list stays under it.
        self.body: list[tuple[str, int, str]] = []
        self.quote = ""
        self.code: list[str] = []
        self.image: str = ""
        self.image_mode = ""      # "bg_right" | "block"
        self.image_pct = 40

    @property
    def bullets(self) -> list[tuple[int, str]]:
        """``(indent_level, text)`` of each bullet."""
        return [(level, text) for kind, level, text in self.body if kind == "bullet"]

    @property
    def paras(self) -> list[str]:
        return [text for kind, _level, text in self.body if kind == "para"]

    @property
    def is_empty(self) -> bool:
        return not any(
            (self.h1, self.h2, self.h3, self.body,
             self.quote, self.code, self.image)
        )


def parse_marp(md: str) -> list[Slide]:
    """Split Marp markdown into slides. Frontmatter is dropped; standalone
    ``---`` lines are slide breaks (the same rule Marp itself uses)."""
    md = _FRONTMATTER.sub("", md or "", count=1)
    slides: list[Slide] = []
    cur = Slide()
    in_code = False

    for raw in md.splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            cur.code.append(raw)
            continue

        if stripped == "---":
            if not cur.is_empty:
                slides.append(cur)
            cur = Slide()
            continue

        if _LEAD.search(stripped):
            cur.lead = True
            continue
        if stripped.startswith("<!--"):
            continue  # any other Marp directive: not renderable here

        m = _IMG.search(stripped)
        if m:
            attrs, src = m.group("attrs"), m.group("src").strip()
            bg = _BG_RIGHT.search(attrs)
            cur.image = src
            if bg:
                cur.image_mode = "bg_right"
                cur.image_pct = int(bg.group("pct") or 40)
            else:
                cur.image_mode = "block"
            # A line that is ONLY an image contributes no text.
            if _IMG.sub("", stripped).strip() == "":
                continue

        if stripped.startswith("### "):
            cur.h3 = stripped[4:].strip()
        elif stripped.startswith("## "):
            cur.h2 = stripped[3:].strip()
        elif stripped.startswith("# "):
            cur.h1 = stripped[2:].strip()
        elif stripped.startswith(">"):
            cur.quote = (cur.quote + " " + stripped.lstrip("> ").strip()).strip()
        elif re.match(r"^[-*+]\s+", stripped):
            indent = (len(line) - len(line.lstrip())) // 2
            cur.body.append(("bullet", min(indent, 1), re.sub(r"^[-*+]\s+", "", stripped)))
        elif stripped:
            cur.body.append(("para", 0, stripped))

    if not cur.is_empty:
        slides.append(cur)
    return slides


# ---------------------------------------------------------------- drawing

# **bold**, *italic*, _italic_ and `code`. As in Markdown, an underscore marks
# emphasis only at a word's edge, so forward_euler and __init__ stay as written.
_INLINE_MARK_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|(?<!\w)_(?=[^_\s])[^_]+?(?<=\S)_(?!\w))")


def _inline_runs(p, text: str, size: float, color, font: str = SANS) -> None:
    """Write ``text`` into paragraph ``p``, honouring **bold**, *italic* and
    _italic_.

    fi.css renders ``strong`` in the darker teal and ``em`` in muted grey; we
    mirror that so emphasis carries the same meaning as in the HTML deck.
    ``$...$`` math inside any of them, but not inside code, becomes a native
    equation (``generation/_pptx_math.py``) in the same colour.
    """
    from pptx.util import Pt
    for i, part in enumerate(_INLINE_MARK_RE.split(text)):
        if not part:
            continue
        marked = i % 2 == 1  # the split puts each marked span at an odd index
        if marked and part.startswith("`"):
            run = p.add_run()
            run.text = part[1:-1]
            run.font.name = MONO
            run.font.color.rgb = _rgb(ACCENT_2)
            run.font.size = Pt(size - 1)
            continue
        bold = italic = False
        part_color = color
        if marked and part.startswith("**"):
            part, bold, part_color = part[2:-2], True, _rgb(ACCENT_2)
        elif marked:
            part, italic, part_color = part[1:-1], True, _rgb(MUTED)
        for segment, is_math in split_math(part):
            if is_math:
                _append_math(p, segment, size, part_color, font, bold=bold)
                continue
            run = p.add_run()
            run.text = segment
            run.font.bold = True if bold else None
            run.font.italic = True if italic else None
            run.font.color.rgb = part_color
            run.font.size = Pt(size)
            run.font.name = font


def _append_math(p, latex: str, size: float, color, font: str, *, bold: bool) -> None:
    """Add one formula to paragraph ``p`` after its runs so far."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
    end = p._p.find(qn("a:endParaRPr"))
    for xml in math_xml(latex, size_pt=size, color=str(color), font=font, bold=bold):
        element = parse_xml(xml)
        if end is not None:
            end.addprevious(element)
        else:
            p._p.append(element)


def _estimate_lines(text: str, width_in: float, size_pt: float,
                    *, wide_factor: float = 0.50) -> int:
    """How many lines ``text`` will wrap to in a ``width_in`` box.

    python-pptx cannot measure text (no font metrics without a rendering
    engine), so this approximates: average glyph advance is roughly
    ``wide_factor`` x the point size. Serif display faces at large sizes run
    wider than body sans, hence the caller-tunable factor. Used only to
    reserve vertical space, so erring one line high is harmless while erring
    low causes a visible overlap.
    """
    if not text:
        return 1
    char_w_in = (size_pt * wide_factor) / 72.0
    per_line = max(8, int(width_in / char_w_in))
    return max(1, -(-len(text) // per_line))  # ceil division


# Body type sizes at full scale: paragraph, level-0 bullet, nested bullet.
_PARA_PT, _BULLET_PT, _SUB_BULLET_PT = 17.0, 16.5, 15.0
# Smallest body scale, about 9 pt: past this a slide is better cut than read.
_MIN_BODY_SCALE = 0.55


# The deck font's single line height, as a multiple of its size. A paragraph's
# line spacing multiplies it: LibreOffice and PowerPoint both advance a 16.5 pt
# bullet at spacing 1.30 by 25.7 pt, where the estimate used 1.30 x 16.5, and a
# figure under a slide's text landed 0.1 pt below its last bullet.
_SINGLE_LINE = 1.2


def _body_height(s: "Slide", width_in: float, k: float) -> float:
    """Estimated height (inches) of a slide's paragraphs and bullets at body
    scale ``k``, by the wrap estimate. Measured against both renderers on
    seven slides, it is within 3.3 pt of LibreOffice and up to 6 pt over
    PowerPoint."""
    # Formulas are estimated by their fallback text, which is about their
    # width on the slide; the LaTeX is several times longer.
    h = space = 0.0
    for kind, level, text in s.body:
        if kind == "para":
            size, space = _PARA_PT * k, 9 * k / 72
            h += _estimate_lines(plain_text(text), width_in, size) * size * 1.32 * _SINGLE_LINE / 72 + space
            continue
        size, space = (_BULLET_PT if level == 0 else _SUB_BULLET_PT) * k, 8 * k / 72
        h += _estimate_lines("•  " + plain_text(text), width_in, size) * size * 1.30 * _SINGLE_LINE / 72 + space
    # Nothing is drawn in the space after the last paragraph.
    return h - space


def _body_scale(s: "Slide", width_in: float, avail_h: float) -> float:
    """The factor (at most 1) that fits a slide's paragraphs and bullets into
    ``avail_h`` by the wrap estimate. A twelve-entry References or Further
    reading slide at full size otherwise ran inches past the slide bottom."""
    k = 1.0
    while k > _MIN_BODY_SCALE and _body_height(s, width_in, k) > avail_h:
        k = round(k - 0.05, 2)
    return max(k, _MIN_BODY_SCALE)


# A figure under a slide's text keeps at least this much height; the text
# shrinks before the figure does. The gap separates the two.
_FIG_MIN_IN, _FIG_GAP_IN = 2.2, 0.2


def _textbox(slide, x, y, w, h):
    from pptx.util import Inches
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _rect(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _set_bg(slide, color) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _fit_picture(slide, img: Path, x, y, w, h) -> None:
    """Place ``img`` contained within the box, preserving aspect ratio.

    ``fit`` in the Marp directive means "show the whole chart" -- the deck
    prompt warns that cropping cuts off a chart's title and axes, so we always
    contain rather than cover.
    """
    from pptx.util import Inches
    try:
        from PIL import Image
        with Image.open(img) as im:
            iw, ih = im.size
        scale = min(w / (iw / 96), h / (ih / 96))
        dw, dh = (iw / 96) * scale, (ih / 96) * scale
    except Exception:  # noqa: BLE001 — PIL missing or unreadable image
        dw, dh = w, h
    slide.shapes.add_picture(
        str(img), Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2),
        width=Inches(dw), height=Inches(dh),
    )


def _footer(slide, page: int, *, dark: bool) -> None:
    """Wordmark + page number, matching fi.css's bottom-left brand mark."""
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    col = _rgb(FAINT if not dark else (0x6E, 0x7C, 0x86))
    tf = _textbox(slide, MARGIN_IN, SLIDE_H_IN - 0.52, 5.0, 0.3)
    r = tf.paragraphs[0].add_run()
    r.text = "F R O N T I E R   I N S I G H T"
    r.font.size, r.font.name, r.font.color.rgb = Pt(9), MONO, col
    tf2 = _textbox(slide, SLIDE_W_IN - MARGIN_IN - 1.0, SLIDE_H_IN - 0.52, 1.0, 0.3)
    tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
    r2 = tf2.paragraphs[0].add_run()
    r2.text = f"{page:02d}"
    r2.font.size, r2.font.name, r2.font.color.rgb = Pt(9), MONO, col


def _render_lead(slide, s: Slide, page: int) -> None:
    """Title / closing slide: inverted ground, the thesis set large in serif.

    fi.css centres these (``_class: lead``) and they carry the 7px accent bar
    along the top edge.
    """
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    _set_bg(slide, INK)
    _rect(slide, 0, 0, SLIDE_W_IN, 0.10, ACCENT)

    body_w = SLIDE_W_IN - 2 * MARGIN_IN
    tf = _textbox(slide, MARGIN_IN, 2.35, body_w, 2.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.12
    _inline_runs(p, s.h1 or s.h2, 40, _rgb(PAPER), font=SERIF)
    for r in p.runs:
        r.font.bold = True
        r.font.color.rgb = _rgb(PAPER)

    if s.h1 and s.h2:
        tf2 = _textbox(slide, MARGIN_IN, 4.95, body_w, 0.9)
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        _inline_runs(p2, s.h2, 18, _rgb((0xC3, 0xCD, 0xD3)), font=SANS)
        for r in p2.runs:
            r.font.color.rgb = _rgb((0xC3, 0xCD, 0xD3))
    if s.paras:
        # The author line the generator puts under the title, or a short
        # line the model wrote on a closing slide. Two lines at most, so a
        # long paragraph cannot run into the footer.
        tf3 = _textbox(slide, MARGIN_IN, 5.85 if (s.h1 and s.h2) else 4.95, body_w, 0.9)
        for i, text in enumerate(s.paras[:2]):
            p3 = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
            p3.alignment = PP_ALIGN.CENTER
            _inline_runs(p3, text, 14, _rgb((0xC3, 0xCD, 0xD3)), font=SANS)
            for r in p3.runs:
                r.font.color.rgb = _rgb((0xC3, 0xCD, 0xD3))
    _footer(slide, page, dark=True)


def _render_content(slide, s: Slide, page: int, figures_dir: Path | None) -> None:
    from pptx.util import Pt
    _set_bg(slide, PAPER)

    has_side_fig = bool(s.image and s.image_mode == "bg_right")
    text_w = SLIDE_W_IN - 2 * MARGIN_IN
    if has_side_fig:
        pct = max(25, min(60, s.image_pct))
        fig_w = SLIDE_W_IN * pct / 100.0
        text_w = SLIDE_W_IN - MARGIN_IN - fig_w - 0.35
        img = _resolve_image(s.image, figures_dir)
        if img:
            _fit_picture(
                slide, img, SLIDE_W_IN - fig_w, 1.0, fig_w - 0.35,
                SLIDE_H_IN - 2.0,
            )

    y = 0.95
    if s.h3:
        tf = _textbox(slide, MARGIN_IN, y, text_w, 0.34)
        r = tf.paragraphs[0].add_run()
        r.text = s.h3.upper()
        r.font.size, r.font.name, r.font.bold = Pt(12), SANS, True
        r.font.color.rgb = _rgb(ACCENT_2)
        y += 0.42

    if s.h2 or s.h1:
        # The teal kicker rule above the title (fi.css h2::before).
        _rect(slide, MARGIN_IN, y, 0.58, 0.042, ACCENT)
        y += 0.20
        title = s.h2 or s.h1
        # Reserve height for the ACTUAL wrapped title. A fixed box let a
        # two-line title overrun into the bullets -- the most visible defect
        # this renderer can produce, and common, since the deck prompt asks
        # for full-sentence findings as slide titles.
        title_pt = 30 if len(title) <= 90 else 26
        lines = _estimate_lines(plain_text(title), text_w, title_pt, wide_factor=0.52)
        line_h = title_pt * 1.12 / 72.0
        title_h = lines * line_h
        tf = _textbox(slide, MARGIN_IN, y, text_w, title_h + 0.1)
        p = tf.paragraphs[0]
        p.line_spacing = 1.12
        _inline_runs(p, title, title_pt, _rgb(INK), font=SERIF)
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = _rgb(INK)
        y += title_h + 0.30

    body_top = max(y, 2.15)
    avail_h = SLIDE_H_IN - body_top - 0.85
    block_img = (
        _resolve_image(s.image, figures_dir) if s.image and s.image_mode == "block" else None
    )
    body_h = 0.0
    if s.body:
        # A figure goes under the text, so the text fits into what is left
        # above the figure's minimum height.
        k = _body_scale(s, text_w, avail_h - (_FIG_MIN_IN + _FIG_GAP_IN if block_img else 0.0))
        body_h = _body_height(s, text_w, k)
        tf = _textbox(slide, MARGIN_IN, body_top, text_w, min(body_h, avail_h) if block_img else avail_h)
        for i, (kind, level, text) in enumerate(s.body):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if kind == "para":
                p.line_spacing = 1.32
                p.space_after = Pt(9 * k)
                _inline_runs(p, text, _PARA_PT * k, _rgb(INK))
                continue
            p.line_spacing = 1.30
            p.space_after = Pt(8 * k)
            marker = p.add_run()
            marker.text = ("    " * level) + ("•  " if level == 0 else "›  ")
            marker.font.size = Pt((16 if level == 0 else 14) * k)
            marker.font.color.rgb = _rgb(ACCENT)   # fi.css li::marker
            marker.font.name = SANS
            _inline_runs(p, text, (_BULLET_PT if level == 0 else _SUB_BULLET_PT) * k, _rgb(INK))

    if s.quote:
        qy = SLIDE_H_IN - 1.9
        _rect(slide, MARGIN_IN, qy, 0.035, 0.72, ACCENT)
        tf = _textbox(slide, MARGIN_IN + 0.28, qy, text_w - 0.28, 0.8)
        p = tf.paragraphs[0]
        p.line_spacing = 1.3
        _inline_runs(p, s.quote, 16, _rgb(MUTED), font=SERIF)
        for r in p.runs:
            r.font.italic = True

    if s.code:
        cy = SLIDE_H_IN - 1.95
        _rect(slide, MARGIN_IN, cy, text_w, 1.15, INK)
        tf = _textbox(slide, MARGIN_IN + 0.2, cy + 0.13, text_w - 0.4, 0.9)
        for i, line in enumerate(s.code[:6]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.25
            r = p.add_run()
            r.text = line[:110]
            r.font.size, r.font.name = Pt(11), MONO
            r.font.color.rgb = _rgb((0xE6, 0xE9, 0xEE))

    if block_img:
        # Under the text's estimated height, not at a fixed offset: a fixed
        # 1.6 in put the figure over the third bullet of a slide with a lead
        # line and three bullets.
        top = body_top + body_h + _FIG_GAP_IN if body_h else body_top
        _fit_picture(
            slide, block_img, MARGIN_IN, top, text_w,
            max(1.5, SLIDE_H_IN - top - 0.85),
        )

    _footer(slide, page, dark=False)


def _resolve_image(src: str, figures_dir: Path | None) -> Path | None:
    """Resolve a deck image reference to a real file, or ``None``.

    Only local paths are resolved: remote URLs cannot be embedded without a
    fetch, and the deck prompt already restricts images to ``figures/``.
    """
    if not src or src.startswith(("http://", "https://", "data:")):
        return None
    p = Path(src)
    if p.is_file():
        return p
    if figures_dir:
        for cand in (figures_dir / p.name, figures_dir.parent / src):
            if cand.is_file():
                return cand
    return None


def render_marp_to_pptx(
    slides_md: Path, out_path: Path, *, figures_dir: Path | None = None,
) -> bool:
    """Render ``slides_md`` to ``out_path``. Returns True on success.

    Never raises: slides are a secondary artifact and a rendering failure must
    not take down a quest that already produced its paper.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as e:  # noqa: BLE001 — python-pptx absent
        _log.warning("python-pptx unavailable (%s); slides.pptx skipped", e)
        return False

    try:
        text = slides_md.read_text(encoding="utf-8")
        parsed = parse_marp(text)
        if not parsed:
            _log.warning("slides.md parsed to zero slides; slides.pptx skipped")
            return False

        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W_IN)
        prs.slide_height = Inches(SLIDE_H_IN)
        blank = prs.slide_layouts[6]

        for i, s in enumerate(parsed, 1):
            slide = prs.slides.add_slide(blank)
            # A slide with only an H1 is a title even without the explicit
            # directive -- the model sometimes omits it on the closing slide.
            if s.lead or (s.h1 and not s.h2 and not s.bullets and not s.paras):
                _render_lead(slide, s, i)
            else:
                _render_content(slide, s, i, figures_dir)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        _log.info(
            "[slides] rendered %d slide(s) to %s via python-pptx (FI theme)",
            len(parsed), out_path.name,
        )
        return True
    except Exception as e:  # noqa: BLE001 — best effort, never fail the quest
        _log.warning("slides.pptx render failed: %r", e)
        return False

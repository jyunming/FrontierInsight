"""The author line (``output.author`` / ``affiliation`` / ``contact_email`` /
``url``) on the rendered outputs: every paper template, the LaTeX and HTML
paper paths, the slides' title slide and the pptx title slide."""

from __future__ import annotations

import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest

from core.config import (
    Config,
    EngineConfig,
    ExecutionConfig,
    KnowledgeConfig,
    OutputConfig,
    ProviderConfig,
)
from core.engine import QuestArtifacts
from generation import paper as paper_mod
from generation._pandoc import find_pandoc
from generation.paper import PaperGenerator, _author_metadata_args
from generation.slides import SlideGenerator, _with_author_line

REPO = Path(__file__).resolve().parent.parent
TEMPLATES = sorted((REPO / "templates" / "paper").glob("*/template.tex"))
AUTHOR = {
    "author": "Jane Chen",
    "affiliation": "R&D_Lab #1",
    "contact_email": "jane_doe@example.org",
    "url": "https://example.org/p",
}


def _output(**fields) -> OutputConfig:  # noqa: ANN003
    return OutputConfig(**fields)


# ---------------------------------------------------------------------------
# Paper


def test_author_metadata_args_skip_unset_fields() -> None:
    assert _author_metadata_args(_output()) == []
    assert _author_metadata_args(_output(author="Jane Chen", url="https://example.org/p")) == [
        "-M", "fi-author=Jane Chen", "-M", "fi-url=https://example.org/p",
    ]


def _render_template(template: Path, tmp_path: Path, args: list[str]) -> str:
    pandoc = find_pandoc(REPO)
    if pandoc is None:
        pytest.skip("pandoc not available")
    md = tmp_path / "in.md"
    md.write_text("Body text.\n", encoding="utf-8")
    done = subprocess.run(
        [pandoc, str(md), "-t", "latex", "--template", str(template),
         "-M", "title=Probe", *args],
        capture_output=True, text=True, encoding="utf-8", timeout=60,
    )
    assert done.returncode == 0, done.stderr
    # Pandoc wraps long lines; a newline is a space to LaTeX.
    return " ".join(done.stdout.split())


@pytest.mark.parametrize("template", TEMPLATES, ids=lambda p: p.parent.name)
def test_every_paper_template_prints_the_escaped_author_line(template: Path, tmp_path: Path) -> None:
    fmt = template.parent.name
    tex = _render_template(template, tmp_path, _author_metadata_args(_output(**AUTHOR)))
    # Pandoc escapes metadata for LaTeX, so & _ # cannot break the compile.
    assert r"R\&D\_Lab \#1" in tex
    assert r"jane\_doe@example.org" in tex
    if fmt == "policy_brief":
        assert r"Jane Chen, R\&D\_Lab \#1, jane\_doe@example.org, https://example.org/p" in tex
    else:
        assert r"\author{Jane Chen\\{\small R\&D\_Lab \#1}\\{\small jane\_doe@example.org}" in tex
    if "pdfauthor" in template.read_text(encoding="utf-8"):
        assert "pdfauthor={Jane Chen}" in tex
    if fmt == "whitepaper":
        assert r"{\large Jane Chen \par}" in tex


@pytest.mark.parametrize("template", TEMPLATES, ids=lambda p: p.parent.name)
def test_every_paper_template_keeps_the_house_byline_without_an_author(template: Path, tmp_path: Path) -> None:
    tex = _render_template(template, tmp_path, [])
    if template.parent.name == "policy_brief":
        assert "Frontier Insight —" in tex
    else:
        assert r"\author{Frontier Insight}" in tex
    assert "fi-author" not in tex


def test_latex_compile_passes_the_author_line_to_pandoc(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    cfg = Config.model_validate({
        "topic": "t", "title": "t",
        "output": {"kinds": ["paper_md", "paper_pdf"],
                   "output_dir": str(tmp_path / "outputs"), **AUTHOR},
    })
    gen = PaperGenerator(cfg)
    monkeypatch.setattr(
        paper_mod.shutil, "which", lambda n: "/fake/pandoc" if n == "pandoc" else None)
    monkeypatch.setattr(gen, "_find_pdf_engine", lambda: ("pdflatex", "/fake/pdflatex"))
    captured: list[str] = []

    def fake_run(cmd, **_kw):  # noqa: ANN001
        captured[:] = list(cmd)
        Path(cmd[cmd.index("-o") + 1]).write_bytes(b"%PDF-1.4 fake\n")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(paper_mod.subprocess, "run", fake_run)
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    paper_md = tmp_path / "paper.md"
    paper_md.write_text("# T\n\nbody\n", encoding="utf-8")

    pdf, skip = gen._compile_pdf(paper_md, out_dir)
    assert skip is None and pdf is not None
    pairs = list(zip(captured, captured[1:]))
    for key, value in (("fi-author", "Jane Chen"), ("fi-affiliation", "R&D_Lab #1"),
                       ("fi-email", "jane_doe@example.org"), ("fi-url", "https://example.org/p")):
        assert ("-M", f"{key}={value}") in pairs


@pytest.mark.slow
def test_a_real_latex_compile_prints_special_characters_in_the_author_line(tmp_path: Path) -> None:
    from generation._pdf_engine import find_pdf_engine
    from generation._pdf_measure import measure_pdf

    if find_pandoc(REPO) is None or find_pdf_engine(REPO) is None:
        pytest.skip("needs pandoc and a LaTeX engine")
    cfg = Config.model_validate({
        "topic": "t", "title": "t",
        "output": {"kinds": ["paper_md", "paper_pdf"], "html_pdf_fallback": False,
                   "output_dir": str(tmp_path / "outputs"), **AUTHOR},
    })
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    paper_md = tmp_path / "paper.md"
    paper_md.write_text("# Probe Title\n\n## Introduction\n\nBody text.\n", encoding="utf-8")

    pdf, skip = PaperGenerator(cfg)._compile_pdf(paper_md, out_dir)
    assert pdf is not None, skip
    doc = measure_pdf(pdf, max_pages=1)
    text = " ".join(line.text for line in doc.pages[0].lines)
    # The OT1 font draws \_ as a rule rather than a glyph, so the PDF's text
    # layer reads it as a space although the page shows an underscore.
    for value in AUTHOR.values():
        assert value.replace("_", " ") in text


# ---------------------------------------------------------------------------
# Slides

DECK = (
    "---\nmarp: true\ntheme: fi\n---\n\n"
    "<!-- _class: lead -->\n\n# The main thesis\n\n## A short kicker\n\n"
    "---\n\n## Finding one\n\n- a point\n\n"
    "---\n\n<!-- _class: lead -->\n\n# Closing line\n"
)


def test_with_author_line_puts_the_line_on_the_title_slide_only() -> None:
    out = _with_author_line(DECK, _output(author="Jane Chen", contact_email="jane@example.org"))
    title, rest = out.split("\n---\n\n## Finding one", 1)
    assert title.rstrip().endswith("## A short kicker\n\nJane Chen · jane@example.org")
    assert "Jane Chen" not in rest


def test_with_author_line_leaves_the_deck_alone_without_an_author() -> None:
    assert _with_author_line(DECK, _output()) == DECK


def test_with_author_line_handles_a_single_slide_deck() -> None:
    out = _with_author_line("# Only slide\n", _output(affiliation="R&D Lab"))
    assert out == "# Only slide\n\nR&D Lab\n"


def test_pptx_title_slide_shows_the_author_line(tmp_path: Path) -> None:
    from pptx import Presentation

    from generation._pptx_slides import render_marp_to_pptx

    md = tmp_path / "slides.md"
    md.write_text(_with_author_line(DECK, _output(**AUTHOR)), encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out) is True
    first = Presentation(str(out)).slides[0]
    texts = [shape.text_frame.text for shape in first.shapes if shape.has_text_frame]
    assert "Jane Chen · R&D_Lab #1 · jane_doe@example.org · https://example.org/p" in texts


@pytest.mark.asyncio
async def test_slide_generator_adds_the_configured_author_line(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    quest_root = tmp_path / "quest"
    (quest_root / "paper").mkdir(parents=True)
    paper_md = quest_root / "paper" / "paper.md"
    paper_md.write_text("# Toy Title\n\nMethods. Results.\n", encoding="utf-8")
    art = QuestArtifacts(quest_id="qid", quest_root=quest_root, paper_md=paper_md, figures_dir=None)
    cfg = Config(
        topic="author line", title="author-line",
        provider=ProviderConfig(name="openai"),
        engine=EngineConfig(max_iterations=1, review_loop=False),
        execution=ExecutionConfig(sandbox="venv", timeout_s=60),
        knowledge=KnowledgeConfig(enabled=False),
        output=OutputConfig(kinds=["slides"], output_dir=tmp_path / "outputs",
                            author="Jane Chen", url="https://example.org/p"),
    )

    async def fake_chat(self, messages, **kw):  # noqa: ANN001
        return DECK

    monkeypatch.setattr("core.provider.LLMClient.chat", fake_chat)
    monkeypatch.setattr("generation.slides.shutil.which", lambda _n: None)

    await SlideGenerator(cfg).generate(art, quest_root)

    body = (quest_root / "slides.md").read_text(encoding="utf-8")
    title_slide = body.split("\n---\n\n## Finding one", 1)[0]
    assert "Jane Chen · https://example.org/p" in title_slide

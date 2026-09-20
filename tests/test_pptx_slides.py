"""Native .pptx slide rendering (``generation/_pptx_slides.py``).

Replaces a `pandoc slides.md -o slides.pptx` shell-out that (a) needed pandoc
on PATH and (b) produced an unthemed deck. These cover the Marp subset the
deck prompt instructs the model to emit.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from generation._pptx_slides import (
    _estimate_lines,
    parse_marp,
    render_marp_to_pptx,
)

DECK = """---
marp: true
theme: fi
---

<!-- _class: lead -->

# The main thesis

## A short kicker

---

### EVIDENCE

## Global error separates the methods

- **RK4** reaches 2.1e-6
- *Verlet* holds 4.8e-4
  - nested detail

> A pull quote.

---

## Chart slide

![bg right:42% fit](figures/x.png)

- point one

---

<!-- _class: lead -->

# Closing line
"""


def test_frontmatter_is_not_a_slide() -> None:
    """The leading `---\nmarp: true\n---` block must not become slide 1."""
    slides = parse_marp(DECK)
    assert len(slides) == 4
    assert slides[0].h1 == "The main thesis"


def test_lead_directive_and_headings() -> None:
    s = parse_marp(DECK)
    assert s[0].lead is True and s[0].h2 == "A short kicker"
    assert s[1].lead is False
    assert s[1].h3 == "EVIDENCE"
    assert s[1].h2 == "Global error separates the methods"
    assert s[3].lead is True


def test_bullets_nesting_and_quote() -> None:
    s = parse_marp(DECK)[1]
    assert [lvl for lvl, _ in s.bullets] == [0, 0, 1]
    assert s.bullets[0][1] == "**RK4** reaches 2.1e-6"
    assert s.quote == "A pull quote."


def test_bg_right_image_parsed_with_percentage() -> None:
    s = parse_marp(DECK)[2]
    assert s.image == "figures/x.png"
    assert s.image_mode == "bg_right"
    assert s.image_pct == 42
    # An image-only line contributes no body text.
    assert s.paras == []
    assert len(s.bullets) == 1


def test_plain_image_is_block_mode() -> None:
    s = parse_marp("## T\n\n![w:900](figures/y.png)\n")[0]
    assert s.image_mode == "block"


def test_directives_other_than_lead_are_dropped() -> None:
    s = parse_marp("<!-- paginate: false -->\n\n## T\n\n- a\n")[0]
    assert s.h2 == "T" and len(s.bullets) == 1


SOURCE_SLIDE = (
    "---\nmarp: true\ntheme: fi\n---\n\n## References\n\n- [1] one source\n- [2] two sources\n\n"
    "_(8 more sources in the paper)_\n"
)


def _body_paragraphs(path: Path) -> list:
    from pptx import Presentation
    return [
        p for sh in Presentation(str(path)).slides[0].shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs if "source" in p.text
    ]


def test_a_line_under_a_list_stays_under_it(tmp_path: Path) -> None:
    """The References slide's "(8 more sources in the paper)" line was drawn
    above its list, and with its underscores."""
    assert [kind for kind, _level, _text in parse_marp(SOURCE_SLIDE)[0].body] == ["bullet", "bullet", "para"]
    md = tmp_path / "slides.md"
    md.write_text(SOURCE_SLIDE, encoding="utf-8")
    assert render_marp_to_pptx(md, tmp_path / "slides.pptx") is True
    paragraphs = _body_paragraphs(tmp_path / "slides.pptx")
    assert [p.text.lstrip("•  ") for p in paragraphs] == [
        "[1] one source", "[2] two sources", "(8 more sources in the paper)",
    ]
    assert all(r.font.italic for r in paragraphs[-1].runs)


def test_an_underscore_inside_a_word_is_not_emphasis(tmp_path: Path) -> None:
    md = tmp_path / "slides.md"
    md.write_text("## T\n\n- forward_euler and __init__ source _stay_ as written\n", encoding="utf-8")
    assert render_marp_to_pptx(md, tmp_path / "slides.pptx") is True
    (p,) = _body_paragraphs(tmp_path / "slides.pptx")
    assert p.text.endswith("forward_euler and __init__ source stay as written")
    assert [r.text for r in p.runs if r.font.italic] == ["stay"]


def test_estimate_lines_detects_wrapping() -> None:
    """Reserving too little height is what made titles overlap the bullets."""
    assert _estimate_lines("Short", 11.5, 30) == 1
    long_title = "Global error separates the integrators by two orders of magnitude"
    assert _estimate_lines(long_title, 11.5, 30, wide_factor=0.52) >= 2
    # Same title in a narrow two-column body wraps further.
    assert (_estimate_lines(long_title, 6.9, 30, wide_factor=0.52)
            > _estimate_lines(long_title, 11.5, 30, wide_factor=0.52))


def test_a_long_source_list_shrinks_to_fit_the_slide(tmp_path: Path) -> None:
    """A twelve-entry References slide at full size ran inches past the slide
    bottom; the body now shrinks until the wrap estimate fits, and a short
    slide keeps full size."""
    entry = (
        "Robert I. McLachlan, G. Quispel, Nicolas Robidoux (1999). Geometric "
        "integration using discrete gradients. Philosophical Transactions of "
        "the Royal Society A DOI: 10.1098/rsta.1999.0363"
    )
    deck = (
        "---\nmarp: true\ntheme: fi\n---\n\n## Short\n\n- one point\n\n---\n\n"
        "## References\n\n"
        + "\n".join(f"- [{i}] {entry}" for i in range(1, 13)) + "\n"
    )
    md = tmp_path / "slides.md"
    md.write_text(deck, encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out) is True

    from pptx import Presentation
    prs = Presentation(str(out))

    def sizes(slide, needle: str) -> set[float]:
        return {
            r.font.size.pt
            for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs if needle in p.text
            for r in p.runs if r.font.size is not None
        }

    assert max(sizes(prs.slides[0], "one point")) == 16.5
    long_list = sizes(prs.slides[1], "DOI")
    assert long_list and max(long_list) < 16.5


def test_renders_a_real_pptx(tmp_path: Path) -> None:
    md = tmp_path / "slides.md"
    md.write_text(DECK, encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out) is True
    assert out.is_file() and out.stat().st_size > 5000

    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 4
    # 16:9 at the same geometry the HTML deck uses.
    assert round(prs.slide_width / 914400, 2) == 13.33
    texts = [
        sh.text_frame.text
        for sh in prs.slides[1].shapes if sh.has_text_frame
    ]
    joined = " ".join(texts)
    assert "Global error separates the methods" in joined
    assert "EVIDENCE" in joined
    assert "RK4" in joined          # bold markers stripped, content kept
    assert "**" not in joined       # markup must not leak into the deck


def test_render_never_raises_on_bad_input(tmp_path: Path) -> None:
    """Slides are secondary: a render failure must not fail the quest."""
    missing = tmp_path / "nope.md"
    assert render_marp_to_pptx(missing, tmp_path / "o.pptx") is False


def test_empty_deck_returns_false(tmp_path: Path) -> None:
    md = tmp_path / "slides.md"
    md.write_text("---\nmarp: true\n---\n", encoding="utf-8")
    assert render_marp_to_pptx(md, tmp_path / "o.pptx") is False


# Slide 5 of the validation quest's deck: a lead line, three bullets (one
# wrapping) and a figure under them. The figure used to sit a fixed 1.6 in
# under the body top, on the third bullet.
FIGURE_DECK = """---
marp: true
theme: fi
---

## Numerical schemes successfully capture exponential decay

**Both RK4 and Velocity-Verlet accurately track the system's dissipative energy loss.**

- The system follows an analytical decay curve for the energy of the damped oscillator.
- Figure 2 shows that both stable methods avoid the catastrophic energy growth seen in Forward Euler.
- The aggregate relative energy error remains low for both methods over the simulation interval.
- Velocity-Verlet keeps its error bounded while Forward Euler grows without limit.

![w:800](figures/energy.png)

---

## A figure alone

![w:800](figures/energy.png)

---

## Many points above a figure

- Point one with enough words to fill most of a line on this slide here.
- Point two with enough words to fill most of a line on this slide here.
- Point three with enough words to fill most of a line on this slide here.
- Point four with enough words to fill most of a line on this slide here.
- Point five with enough words to fill most of a line on this slide here.
- Point six with enough words to fill most of a line on this slide here.
- Point seven with enough words to fill most of a line on this slide here.

![w:800](figures/energy.png)
"""


def _figure_deck(tmp_path: Path) -> Path:
    from PIL import Image, ImageDraw

    figures = tmp_path / "figures"
    figures.mkdir()
    chart = Image.new("RGB", (1500, 900), "white")
    ImageDraw.Draw(chart).rectangle((40, 40, 1460, 860), outline="black", width=8)
    chart.save(figures / "energy.png")
    md = tmp_path / "slides.md"
    md.write_text(FIGURE_DECK, encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out, figures_dir=figures) is True
    return out


def test_a_figure_goes_under_the_text_it_follows(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(_figure_deck(tmp_path)))
    slide = prs.slides[0]
    picture = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    body = next(sh for sh in slide.shapes if sh.has_text_frame and "aggregate relative" in sh.text_frame.text)
    assert picture.top >= body.top + body.height
    assert picture.top + picture.height <= prs.slide_height - int(0.85 * 914400)
    assert picture.height >= int(2.0 * 914400), "the figure stays readable"
    # With no text above it the figure starts where the body would.
    alone = next(sh for sh in prs.slides[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert alone.top < picture.top
    # Seven points leave no room at full size: the text shrinks, and the
    # figure keeps its minimum height above the footer.
    crowded = next(sh for sh in prs.slides[2].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert crowded.top + crowded.height <= prs.slide_height - int(0.85 * 914400)
    assert crowded.height >= int(2.0 * 914400)


def test_a_figure_slide_gives_its_figure_the_slide_width_and_a_slide_with_bullets_keeps_the_text_width(
    tmp_path: Path,
) -> None:
    """A figure is drawn at its box's width and its tick labels shrink with it: a slide
    of a title, a lead line and the figure gives it the room out to 0.75 in, one that also
    has bullets keeps the text's margins. Under 90% of the slide's width, past which the
    measurements of a rendered slide would take the figure for a full-bleed decoration."""
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from generation._pptx_slides import MARGIN_IN, SLIDE_W_IN, _FIG_SLIDE_MARGIN_IN

    (tmp_path / "figures").mkdir()
    Image.new("RGB", (1500, 450), "white").save(tmp_path / "figures" / "row.png", dpi=(100, 100))
    md = tmp_path / "slides.md"
    md.write_text(
        "## A figure slide\n\n**The lead line of the slide.**\n\n![](figures/row.png)\n\n---\n\n"
        "## A slide with a bullet\n\n- One point.\n\n![](figures/row.png)\n",
        encoding="utf-8",
    )
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out, figures_dir=tmp_path / "figures") is True
    alone, shared = (
        next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE) for slide in Presentation(str(out)).slides
    )
    inch = 914400
    assert abs(alone.width / inch - (SLIDE_W_IN - 2 * _FIG_SLIDE_MARGIN_IN)) < 0.01
    assert abs(shared.width / inch - (SLIDE_W_IN - 2 * MARGIN_IN)) < 0.01
    assert alone.width < 0.9 * SLIDE_W_IN * inch
    assert alone.top + alone.height <= 7.5 * inch - int(0.6 * inch) + 1


@pytest.mark.slow
def test_the_exported_deck_has_no_figure_over_its_text(tmp_path: Path) -> None:
    from generation._office_pdf import find_libreoffice, pptx_to_pdf
    from generation._pdf_measure import measure_pdf, slides_report

    if find_libreoffice() is None:
        pytest.skip("LibreOffice is not installed")
    pdf, reason = pptx_to_pdf(_figure_deck(tmp_path), tmp_path / "export")
    assert pdf is not None, reason
    report = slides_report(measure_pdf(pdf))
    assert [f for f in report["findings"] if f["check"] in ("overlap", "overflow", "figure_gap")] == []


def test_the_text_estimate_is_as_tall_as_the_renderers_draw_it() -> None:
    """Five bullets at full size: LibreOffice drew them 158 pt tall and
    PowerPoint 155 pt. The estimate had them at 147 pt, so a figure under
    them landed on the last bullet."""
    from generation._pptx_slides import MARGIN_IN, SLIDE_W_IN, _body_height

    five = parse_marp("## T\n\n" + "".join(f"- Bullet {n} with a few words.\n" for n in range(5)))[0]
    assert 155 <= _body_height(five, SLIDE_W_IN - 2 * MARGIN_IN, 1.0) * 72 <= 166


# The validation quest's formulas, as its deck wrote them.
MATH_DECK = r"""---
marp: true
theme: fi
---

## Forward Euler is unstable at standard time steps

**The damping term $c x'$ alters the trade-off.**

- At $h = 0.01$, RK4 error is $5.297 \times 10^{-9}\text{m}$.
- The decay follows $E_{analytical}(t) = E_0 \exp(-(c/m)t)$.
- Set `$HOME` before running.
"""


def _math_deck(tmp_path: Path) -> Path:
    md = tmp_path / "slides.md"
    md.write_text(MATH_DECK, encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out) is True
    return out


def test_formulas_become_native_equations_in_place(tmp_path: Path) -> None:
    import re
    import zipfile

    xml = zipfile.ZipFile(_math_deck(tmp_path)).read("ppt/slides/slide1.xml").decode("utf-8")
    assert xml.count("<m:oMath") == 4 and xml.count("<mc:Fallback>") == 4
    texts = re.findall(r"<a:t>([^<]*)</a:t>", xml)
    # No LaTeX left in any text run; code stays code.
    assert [t for t in texts if "$" in t or "\\" in t] == ["$HOME"]
    # Text, equation, text: the formula sits where it was written.
    at = xml.index(">At <")
    assert at < xml.index("<mc:AlternateContent", at) < xml.index(">, RK4 error is <")
    # The bold lead line's formula keeps the lead line's colour.
    start = xml.index("<mc:AlternateContent", xml.index("The damping term"))
    lead_formula = xml[start: xml.index("</mc:AlternateContent>", start)]
    assert "<m:oMath" in lead_formula
    assert 'val="0A4F4D"' in lead_formula and 'val="16222B"' not in lead_formula


def test_a_formula_takes_the_room_of_its_fallback_text_not_of_its_latex() -> None:
    """Four formulas are 120 characters of LaTeX but about 50 on the slide;
    estimating the LaTeX would wrap the line and shrink the slide's text."""
    from generation._pptx_slides import _body_height

    formula = r"$5.297 \times 10^{-9}\text{m}$"
    written = parse_marp("## T\n\n- " + " and ".join([formula] * 4) + "\n")[0]
    shown = parse_marp("## T\n\n- " + " and ".join(["5.297 × 10−9 m"] * 4) + "\n")[0]
    assert _body_height(written, 11.5, 1.0) == _body_height(shown, 11.5, 1.0)


@pytest.mark.slow
def test_libreoffice_shows_the_readable_fallback(tmp_path: Path) -> None:
    from generation._office_pdf import find_libreoffice, pptx_to_pdf
    from generation._pdf_measure import measure_pdf, slides_report

    if find_libreoffice() is None:
        pytest.skip("LibreOffice is not installed")
    pdf, reason = pptx_to_pdf(_math_deck(tmp_path), tmp_path / "export")
    assert pdf is not None, reason
    doc = measure_pdf(pdf)
    text = " ".join(line.text for page in doc.pages for line in page.lines)
    assert "×" in text and "\\times" not in text and "$h" not in text
    assert [f for f in slides_report(doc)["findings"] if f["check"] == "raw_markup"] == []


def _powerpoint_free() -> str | None:
    """Why PowerPoint cannot export in this test, or None when it can."""
    import shutil
    import subprocess
    import sys

    if sys.platform != "win32":
        return "PowerPoint export needs Windows"
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return "pywin32 is not installed"
    running = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/NH"], capture_output=True, text=True, check=False,
    ).stdout
    if "POWERPNT.EXE" in running.upper():
        return "PowerPoint is open; the test never touches the user's window"
    if shutil.which("POWERPNT") is None and not any(
        Path(root, "Microsoft Office", "root", "Office16", "POWERPNT.EXE").is_file()
        for root in (r"C:\Program Files", r"C:\Program Files (x86)")
    ):
        return "PowerPoint is not installed"
    return None


@pytest.mark.slow
def test_powerpoint_draws_the_equations(tmp_path: Path) -> None:
    reason = _powerpoint_free()
    if reason:
        pytest.skip(reason)
    import win32com.client

    from generation._pdf_measure import measure_pdf, slides_report

    pptx = _math_deck(tmp_path)
    pdf = tmp_path / "powerpoint.pdf"
    # pytest's faulthandler may print "Windows fatal exception: code
    # 0x800706be" during these COM calls; they return normally and the
    # test's result stands.
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        deck = app.Presentations.Open(str(pptx), -1, 0, 0)   # read-only, no window
        deck.SaveAs(str(pdf), 32)                           # ppSaveAsPDF
        deck.Close()
    finally:
        app.Quit()
    doc = measure_pdf(pdf)
    text = " ".join(line.text for page in doc.pages for line in page.lines)
    assert "$" not in text.replace("$HOME", "") and "\\times" not in text
    # PowerPoint draws the equation, not the fallback: its h is the math
    # italic ℎ (U+210E), where the fallback text has a plain h.
    assert "ℎ = 0.01" in text and "×" in text
    assert [f for f in slides_report(doc)["findings"] if f["check"] == "raw_markup"] == []


# ---------------------------------------------------------------- tables

# The validation quest's slide 4: a bold lead line, a table, two bullets. It
# printed the table as five lines of pipes.
TABLE_DECK = r"""---
marp: true
theme: fi
---

## Stronger transmission makes agreement immediate

**For $R_0=3.0$, large outbreaks follow the deterministic size.**

| Population | Conditional rate | Note |
|---:|:---:|:---|
| 100 | **0.930** | _low_ |
| 1,000 | 0.940 | `code` |
| $N=5{,}000$ | $0.941$ | a $x^2$ b |

- The magnitude is predictable.
- Its occurrence remains probabilistic.
"""


def _table_pptx(tmp_path: Path, deck: str = TABLE_DECK) -> Path:
    md = tmp_path / "slides.md"
    md.write_text(deck, encoding="utf-8")
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out) is True
    return out


def _slide_xml(pptx: Path, number: int = 1) -> str:
    import zipfile

    return zipfile.ZipFile(pptx).read(f"ppt/slides/slide{number}.xml").decode("utf-8")


def _table_shape(pptx: Path):
    from pptx import Presentation

    return next(sh for sh in Presentation(str(pptx)).slides[0].shapes if sh.has_table)


def test_a_table_is_parsed_with_its_cells_and_alignment() -> None:
    slide = parse_marp(TABLE_DECK)[0]
    (table,) = slide.tables
    assert table.header == ["Population", "Conditional rate", "Note"]
    assert table.aligns == ["r", "c", "l"]
    assert [row[0] for row in table.rows] == ["100", "1,000", "$N=5{,}000$"]
    # It keeps its place between the lead line and the bullets, and its lines
    # are no paragraphs.
    assert [kind for kind, _index, _text in slide.body] == ["para", "table", "bullet", "bullet"]
    assert len(slide.paras) == 1 and len(slide.bullets) == 2


def test_table_syntax_variants_are_read() -> None:
    """Every delimiter-row spelling the models wrote, rows short or long of the
    header, and an escaped pipe."""
    for delim in ("|---|---|", "| :--- | :--- |", "|----------|----------|", "|:---:|---:|"):
        (table,) = parse_marp(f"## T\n\n| a | b |\n{delim}\n| 1 | 2 |\n")[0].tables
        assert table.rows == [["1", "2"]], delim
    assert parse_marp("## T\n\n| a | b |\n|:---:|---:|\n| 1 | 2 |\n")[0].tables[0].aligns == ["c", "r"]
    (table,) = parse_marp("## T\n\n| a | b |\n|---|---|\n| only |\n| 1 | 2 | 3 |\n| x \\| y | z |\n")[0].tables
    assert table.rows == [["only", ""], ["1", "2"], ["x | y", "z"]]


def test_pipe_text_that_is_not_a_table_stays_text() -> None:
    """A pipe line needs a delimiter row under it, and a bare --- is still a
    slide break."""
    slide = parse_marp("## T\n\n| not | a table |\n\nafter\n")[0]
    assert slide.tables == [] and slide.paras == ["| not | a table |", "after"]
    slides = parse_marp("## T\n\n| a | b |\n---\n## Next\n\n- x\n")
    assert len(slides) == 2 and slides[0].tables == []


def test_a_table_becomes_a_native_table_not_pipes(tmp_path: Path) -> None:
    import re

    from pptx import Presentation

    pptx = _table_pptx(tmp_path)
    xml = _slide_xml(pptx)
    assert xml.count("<a:tbl>") == 1
    assert [t for t in re.findall(r"<a:t>([^<]*)</a:t>", xml) if "|" in t or "---" in t] == []

    table = _table_shape(pptx).table
    assert (len(table.rows), len(table.columns)) == (4, 3)
    assert [table.cell(0, c).text for c in range(3)] == ["Population", "Conditional rate", "Note"]
    assert [table.cell(r, 0).text for r in (1, 2)] == ["100", "1,000"]
    assert table.cell(1, 1).text == "0.930" and table.cell(2, 2).text == "code"
    # The bullets under it are still text.
    slide = Presentation(str(pptx)).slides[0]
    assert any("magnitude is predictable" in sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_a_column_keeps_the_alignment_the_delimiter_row_gave_it(tmp_path: Path) -> None:
    from pptx.enum.text import PP_ALIGN

    table = _table_shape(_table_pptx(tmp_path)).table
    for row in range(4):
        alignments = [table.cell(row, c).text_frame.paragraphs[0].alignment for c in range(3)]
        assert alignments == [PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.LEFT]


def test_a_formula_in_a_cell_is_a_native_equation(tmp_path: Path) -> None:
    import re

    xml = _slide_xml(_table_pptx(tmp_path))
    frame = xml[xml.index("<a:tbl>"): xml.index("</a:tbl>")]
    # $N=5{,}000$, $0.941$ and the $x^2$ inside a sentence.
    assert len(re.findall(r"<m:oMath[ >]", frame)) == 3 and frame.count("<mc:Fallback>") == 3
    assert [t for t in re.findall(r"<a:t>([^<]*)</a:t>", frame) if "$" in t or "\\" in t] == []
    # Text, equation, text: the formula sits where it was written.
    at = frame.index(">a <")
    assert at < frame.index("<mc:AlternateContent", at) < frame.index("> b<")
    # A cell that is only a formula keeps its column's alignment (PowerPoint
    # would set it in the middle); the one inside a sentence stays inline.
    assert re.findall(r'<m:jc m:val="(\w+)"/>', frame) == ["right", "center"]


def test_a_table_is_drawn_in_the_fi_theme(tmp_path: Path) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_FILL

    pptx = _table_pptx(tmp_path)
    xml = _slide_xml(pptx)
    # Not Office's blue banded default.
    assert "5C22544A" not in xml
    table = _table_shape(pptx).table
    assert table.first_row is True and table.horz_banding is False
    header, body = table.cell(0, 0), table.cell(1, 0)
    assert header.fill.fore_color.rgb == RGBColor(0xEE, 0xF1, 0xEE)
    assert body.fill.type == MSO_FILL.BACKGROUND, "no fill under the body: the paper shows"
    head_run = header.text_frame.paragraphs[0].runs[0]
    assert head_run.font.bold and head_run.font.color.rgb == RGBColor(0x0A, 0x4F, 0x4D)
    assert head_run.font.name == "Segoe UI" and head_run.font.size.pt == 14.5
    body_run = body.text_frame.paragraphs[0].runs[0]
    assert not body_run.font.bold and body_run.font.color.rgb == RGBColor(0x16, 0x22, 0x2B)
    # A 2 pt rule of the theme's teal under the header, 1 pt hairlines under the rows.
    assert 'w="25400"' in xml and 'val="0E6E6B"' in xml and 'w="12700"' in xml and 'val="E4E2DA"' in xml


def test_text_above_and_below_a_table_keeps_its_place(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(_table_pptx(tmp_path)))
    slide = prs.slides[0]
    table = next(sh for sh in slide.shapes if sh.has_table)
    lead = next(sh for sh in slide.shapes if sh.has_text_frame and "large outbreaks" in sh.text_frame.text)
    bullets = next(sh for sh in slide.shapes if sh.has_text_frame and "magnitude is predictable" in sh.text_frame.text)
    assert lead.top + lead.height <= table.top
    assert table.top + table.height <= bullets.top
    assert bullets.top + bullets.height <= prs.slide_height - int(0.85 * 914400)


def test_a_table_takes_the_room_of_the_slide_body(tmp_path: Path) -> None:
    """A table counts in the height the deck's own scale logic fits: fourteen
    rows shrink the table's type from 14.5 pt so it ends above the footer, and
    a table too tall even at its floor stops at 10 pt."""
    from pptx import Presentation

    def built(rows: int):
        body = "".join(f"| row {i} | {i} |\n" for i in range(rows))
        prs = Presentation(str(_table_pptx(tmp_path, f"## Many rows\n\n| Name | Value |\n|---|---|\n{body}")))
        table = next(sh for sh in prs.slides[0].shapes if sh.has_table)
        return table, table.table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt, prs

    _, short_pt, _ = built(4)
    assert short_pt == 14.5
    tall, tall_pt, prs = built(14)
    assert 10 <= tall_pt < 14.5
    assert tall.top + tall.height <= prs.slide_height - int(0.85 * 914400)
    _, floor_pt, _ = built(40)
    assert floor_pt == 10.0


def test_a_table_over_a_figure_leaves_the_figure_its_room(tmp_path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    (tmp_path / "figures").mkdir()
    Image.new("RGB", (1500, 900), "white").save(tmp_path / "figures" / "energy.png")
    md = tmp_path / "slides.md"
    md.write_text(
        "## T\n\n| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n\n![w:800](figures/energy.png)\n", encoding="utf-8",
    )
    out = tmp_path / "slides.pptx"
    assert render_marp_to_pptx(md, out, figures_dir=tmp_path / "figures") is True
    prs = Presentation(str(out))
    table = next(sh for sh in prs.slides[0].shapes if sh.has_table)
    picture = next(sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture.top >= table.top + table.height
    assert picture.height >= int(2.0 * 914400)


def test_a_title_with_only_a_table_is_not_a_title_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    pptx = _table_pptx(tmp_path, "# Results\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    assert any(sh.has_table for sh in Presentation(str(pptx)).slides[0].shapes)


def test_a_tidy_table_is_not_counted_as_wrapping() -> None:
    """The header's longest cell fills its column exactly; a rounding error in
    the wrap estimate once counted it as two lines and drew a header row
    almost twice the height of the rows."""
    from generation._pptx_slides import MARGIN_IN, SLIDE_W_IN, _table_layout

    body_w = SLIDE_W_IN - 2 * MARGIN_IN
    (table,) = parse_marp(
        "## T\n\n| Population | Conditional attack rate | Deterministic attack rate |\n|---:|---:|---:|\n"
        "| 100 | 0.930 | 0.941 |\n| 1,000 | 0.940 | 0.941 |\n"
    )[0].tables
    size, widths, heights = _table_layout(table, body_w, 1.0)
    assert size == 14.5
    assert len({round(h, 4) for h in heights}) == 1, "every row is one line"
    assert sum(widths) < body_w, "a compact table does not fill the slide"


def test_a_wide_table_wraps_its_cells_inside_the_body() -> None:
    from generation._pptx_slides import MARGIN_IN, SLIDE_W_IN, _table_layout

    body_w = SLIDE_W_IN - 2 * MARGIN_IN
    long_cell = "Sauropodomorpha repositioned within the tree by the newer analysis " * 3
    (table,) = parse_marp(f"## T\n\n| Old | New |\n|---|---|\n| short | {long_cell} |\n")[0].tables
    _size, widths, heights = _table_layout(table, body_w, 1.0)
    assert abs(sum(widths) - body_w) < 1e-6
    assert heights[1] > 2 * heights[0], "the long cell wraps onto several lines"


@pytest.mark.slow
def test_libreoffice_draws_the_table_without_pipes_or_overlap(tmp_path: Path) -> None:
    from generation._office_pdf import find_libreoffice, pptx_to_pdf
    from generation._pdf_measure import measure_pdf, slides_report

    if find_libreoffice() is None:
        pytest.skip("LibreOffice is not installed")
    pdf, reason = pptx_to_pdf(_table_pptx(tmp_path), tmp_path / "export")
    assert pdf is not None, reason
    doc = measure_pdf(pdf)
    text = " ".join(line.text for page in doc.pages for line in page.lines)
    assert "|" not in text and "0.930" in text and "Population" in text
    findings = slides_report(doc)["findings"]
    assert [f for f in findings if f["check"] in ("overlap", "overflow", "raw_markup")] == []
